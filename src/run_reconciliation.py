"""Reconcile payment records and generate Excel and CSV reports."""

from __future__ import annotations

import os
import tempfile
from decimal import Decimal
from getpass import getpass
from pathlib import Path
from typing import TYPE_CHECKING, Any

import pandas as pd
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

if TYPE_CHECKING:
    import psycopg


ROOT_DIR = Path(__file__).resolve().parent.parent
REPORTS_DIR = ROOT_DIR / "reports"

DETECTED_QUERY = """
SELECT
    defect_type,
    entity_id,
    expected_amount,
    actual_amount,
    variance_amount,
    amount_under_review,
    reason
FROM public.reconciliation_exceptions
ORDER BY defect_type, entity_id;
"""

GROUND_TRUTH_QUERY = """
SELECT
    entity_type,
    entity_id,
    defect_type,
    expected_amount,
    impact_basis,
    description
FROM public.ground_truth_exceptions
ORDER BY defect_type, entity_id;
"""

ROW_COUNTS_QUERY = """
SELECT 'internal_orders' AS table_name, COUNT(*) AS row_count
FROM public.internal_orders
UNION ALL
SELECT 'gateway_statement', COUNT(*)
FROM public.gateway_statement
UNION ALL
SELECT 'bank_statement', COUNT(*)
FROM public.bank_statement
UNION ALL
SELECT 'ground_truth_exceptions', COUNT(*)
FROM public.ground_truth_exceptions
UNION ALL
SELECT 'reconciliation_exceptions', COUNT(*)
FROM public.reconciliation_exceptions
ORDER BY table_name;
"""

SHEET_NAMES = {
    "MISSING_GATEWAY_RECORD": "Missing_Gateway",
    "ORPHAN_GATEWAY_RECORD": "Orphan_Gateway",
    "INTERNAL_GATEWAY_STATUS_MISMATCH": "Status_Mismatch",
    "ORDER_GATEWAY_AMOUNT_MISMATCH": "Order_Gateway_Amount",
    "MDR_OVERCHARGE": "MDR_Overcharges",
    "GATEWAY_NET_PAYOUT_MISMATCH": "Net_Payout_Mismatch",
    "SETTLEMENT_TIMING_LAG": "Settlement_Lags",
    "DUPLICATE_BANK_CREDIT": "Duplicate_Credits",
    "BANK_SETTLEMENT_AMOUNT_MISMATCH": "Bank_Amount_Mismatch",
    "NON_CREDITED_BANK_ENTRY": "Non_Credited_Entries",
    "MISSING_BANK_SETTLEMENT": "Missing_Settlements",
    "ORPHAN_BANK_CREDIT": "Orphan_Bank_Credits",
}

CURRENCY_COLUMNS = {
    "expected_amount",
    "actual_amount",
    "variance_amount",
    "amount_under_review",
    "amount_difference",
    "gross_amount",
    "fee_deducted",
    "gst",
    "net_payout",
    "credit_amount",
    "actual_credit_total",
    "expected_payout",
    "expected_fee",
    "expected_gst",
    "actual_fee",
    "actual_gst",
}



def get_connection() -> "psycopg.Connection":
    """Connect only when a database target has been explicitly configured."""
    dbname = os.getenv("PGDATABASE", "").strip()
    if not dbname:
        raise ValueError(
            "PGDATABASE is not set. Set it explicitly to the dedicated project "
            "database before running this script; no database default is used."
        )

    import psycopg

    host = os.getenv("PGHOST", "localhost")
    port = int(os.getenv("PGPORT", "5432"))
    user = os.getenv("PGUSER", "postgres")
    print(
        "PostgreSQL target: "
        f"host={host}, port={port}, database={dbname}, user={user}"
    )

    password = os.getenv("PGPASSWORD")
    if not password:
        password = getpass("PostgreSQL password for the configured user: ")

    return psycopg.connect(
        host=host,
        port=port,
        dbname=dbname,
        user=user,
        password=password,
        connect_timeout=10,
    )



def fetch_dataframe(
    conn: "psycopg.Connection",
    query: str,
) -> pd.DataFrame:
    """Run a query and return its results as a DataFrame."""
    with conn.cursor() as cursor:
        cursor.execute(query)
        rows = cursor.fetchall()
        columns = (
            [column.name for column in cursor.description]
            if cursor.description
            else []
        )

    return pd.DataFrame(rows, columns=columns)


def as_excel_safe(df: pd.DataFrame) -> pd.DataFrame:
    """Convert Decimal values to numbers that Excel can display."""
    safe = df.copy()

    for column in safe.columns:
        if safe[column].dtype == "object":
            safe[column] = safe[column].map(
                lambda value: (
                    float(value) if isinstance(value, Decimal) else value
                )
            )

    return safe


def build_validation(
    detected: pd.DataFrame,
    ground_truth: pd.DataFrame,
) -> tuple[pd.DataFrame, pd.DataFrame]:
    """Compare detected exceptions with the synthetic answer key safely."""
    key_columns = ["defect_type", "entity_id"]
    required_detected = set(key_columns + ["amount_under_review"])
    required_expected = set(key_columns + ["expected_amount"])

    missing_detected = required_detected - set(detected.columns)
    missing_expected = required_expected - set(ground_truth.columns)
    if missing_detected or missing_expected:
        raise ValueError(
            "Validation input is missing required columns. "
            f"Detected missing: {sorted(missing_detected)}; "
            f"ground-truth missing: {sorted(missing_expected)}."
        )

    for label, frame in (("detected", detected), ("ground truth", ground_truth)):
        if frame[key_columns].isna().any().any():
            raise ValueError(
                f"The {label} input contains a blank defect_type or entity_id. "
                "Each exception needs a complete validation key."
            )
        duplicates = frame.duplicated(key_columns, keep=False)
        if duplicates.any():
            examples = (
                frame.loc[duplicates, key_columns]
                .drop_duplicates()
                .head(5)
                .to_dict(orient="records")
            )
            raise ValueError(
                f"The {label} input contains duplicate (defect_type, entity_id) "
                f"keys. Examples: {examples}"
            )

    expected = ground_truth[
        ["defect_type", "entity_id", "expected_amount"]
    ].copy()
    found = detected[
        ["defect_type", "entity_id", "amount_under_review"]
    ].copy()

    comparison = expected.merge(
        found,
        on=key_columns,
        how="outer",
        indicator=True,
        validate="one_to_one",
    )

    comparison["match_status"] = comparison["_merge"].map(
        {
            "both": "MATCHED",
            "left_only": "FALSE_NEGATIVE",
            "right_only": "FALSE_POSITIVE",
        }
    )

    def amount_difference(row: pd.Series) -> float | None:
        if row["_merge"] != "both":
            return None
        expected_amount = row["expected_amount"]
        actual_amount = row["amount_under_review"]
        if pd.isna(expected_amount) or pd.isna(actual_amount):
            return None
        return round(float(actual_amount) - float(expected_amount), 2)

    comparison["amount_difference"] = comparison.apply(
        amount_difference, axis=1
    )

    def amount_matches(row: pd.Series) -> bool | None:
        if row["_merge"] != "both":
            return None
        if pd.isna(row["expected_amount"]) or pd.isna(row["amount_under_review"]):
            return False
        return abs(float(row["amount_difference"])) < 0.005

    comparison["amount_match"] = comparison.apply(amount_matches, axis=1)

    categories = sorted(
        set(ground_truth["defect_type"].dropna())
        | set(detected["defect_type"].dropna())
    )
    summary_rows: list[dict[str, Any]] = []

    for category in categories:
        expected_count = int(
            (ground_truth["defect_type"] == category).sum()
        )
        detected_count = int(
            (detected["defect_type"] == category).sum()
        )
        category_comparison = comparison[
            comparison["defect_type"] == category
        ]

        matched = int(
            (category_comparison["match_status"] == "MATCHED").sum()
        )
        false_positives = int(
            (category_comparison["match_status"] == "FALSE_POSITIVE").sum()
        )
        false_negatives = int(
            (category_comparison["match_status"] == "FALSE_NEGATIVE").sum()
        )
        amount_matches_count = int(
            (category_comparison["amount_match"] == True).sum()  # noqa: E712
        )
        amount_mismatches = int(
            (
                (category_comparison["match_status"] == "MATCHED")
                & (category_comparison["amount_match"] == False)  # noqa: E712
            ).sum()
        )

        precision = (
            100.0 * matched / detected_count if detected_count else None
        )
        recall = 100.0 * matched / expected_count if expected_count else None

        summary_rows.append(
            {
                "defect_type": category,
                "expected_cases": expected_count,
                "detected_cases": detected_count,
                "matched_cases": matched,
                "false_positives": false_positives,
                "false_negatives": false_negatives,
                "amount_matches": amount_matches_count,
                "amount_mismatches": amount_mismatches,
                "precision_pct": (
                    round(precision, 2) if precision is not None else None
                ),
                "recall_pct": round(recall, 2) if recall is not None else None,
            }
        )

    summary = pd.DataFrame(summary_rows)
    return summary, comparison.drop(columns=["_merge"])



def style_workbook(path: Path) -> None:
    """Format the workbook so the reports are easier to read."""
    from openpyxl import load_workbook

    workbook = load_workbook(path)
    header_fill = PatternFill(fill_type="solid", fgColor="1F4E78")
    header_font = Font(color="FFFFFF", bold=True)

    for worksheet in workbook.worksheets:
        worksheet.freeze_panes = "A2"

        if worksheet.max_row >= 1 and worksheet.max_column >= 1:
            worksheet.auto_filter.ref = worksheet.dimensions

            for cell in worksheet[1]:
                cell.fill = header_fill
                cell.font = header_font
                cell.alignment = Alignment(vertical="top", wrap_text=True)

        header_to_column = {
            str(cell.value): cell.column
            for cell in worksheet[1]
            if cell.value is not None
        }
        for currency_header in CURRENCY_COLUMNS & set(header_to_column):
            column_index = header_to_column[currency_header]
            for row_index in range(2, worksheet.max_row + 1):
                cell = worksheet.cell(row=row_index, column=column_index)
                if isinstance(cell.value, (int, float)) and not isinstance(cell.value, bool):
                    cell.number_format = "#,##0.00;[Red]-#,##0.00"

        for column_cells in worksheet.columns:
            column_index = column_cells[0].column
            max_length = 0

            for cell in column_cells:
                value = "" if cell.value is None else str(cell.value)
                max_length = max(max_length, min(len(value), 60))
                cell.alignment = Alignment(vertical="top", wrap_text=True)

            worksheet.column_dimensions[
                get_column_letter(column_index)
            ].width = min(max(max_length + 2, 12), 42)

    workbook.save(path)


def main() -> int:
    """Read the database, validate exceptions, and write report files."""
    if not os.getenv("PGDATABASE", "").strip():
        print(
            "PGDATABASE is not set. Set it explicitly to the dedicated project "
            "database before running this script. No connection was attempted."
        )
        return 1

    try:
        import psycopg
    except ImportError:
        print(
            "Missing dependency: psycopg. "
            "Run: python -m pip install -r requirements.txt"
        )
        return 1

    REPORTS_DIR.mkdir(parents=True, exist_ok=True)

    try:
        with get_connection() as conn:
            detected = fetch_dataframe(conn, DETECTED_QUERY)
            ground_truth = fetch_dataframe(conn, GROUND_TRUTH_QUERY)
            row_counts = fetch_dataframe(conn, ROW_COUNTS_QUERY)
    except ValueError as exc:
        print(f"Configuration error: {exc}")
        return 1
    except psycopg.Error as exc:
        print("Database operation failed.")
        print(
            "Check that PostgreSQL is running, the configured database exists, "
            "the password is correct, and the tables and view exist."
        )
        print(f"PostgreSQL detail: {exc}")
        return 1

    if detected.empty:
        print(
            "The reconciliation view returned no exceptions. "
            "Check whether this is expected for the current data."
        )

    if ground_truth.empty:
        print(
            "The ground-truth table is empty. Import "
            "data/test/ground_truth_exceptions.csv before validation."
        )
        return 1

    try:
        summary, comparison = build_validation(detected, ground_truth)
    except (ValueError, KeyError) as exc:
        print(f"Validation cannot proceed: {exc}")
        return 1

    # Use one stable workbook name so repeated runs overwrite the prior report.
    workbook_path = REPORTS_DIR / "payment_reconciliation_report.xlsx"
    exception_csv = REPORTS_DIR / "reconciliation_exceptions.csv"
    validation_csv = REPORTS_DIR / "validation_detail.csv"

    detected_export = as_excel_safe(detected)
    summary_export = as_excel_safe(summary)
    comparison_export = as_excel_safe(comparison)
    ground_truth_export = as_excel_safe(ground_truth)
    row_counts_export = as_excel_safe(row_counts)

    temporary_paths: dict[Path, Path] = {}
    try:
        for final_path in (workbook_path, exception_csv, validation_csv):
            handle = tempfile.NamedTemporaryFile(
                prefix=f".{final_path.stem}.",
                suffix=final_path.suffix,
                dir=REPORTS_DIR,
                delete=False,
            )
            handle.close()
            temporary_paths[final_path] = Path(handle.name)

        detected_export.to_csv(
            temporary_paths[exception_csv],
            index=False,
            encoding="utf-8-sig",
        )
        comparison_export.to_csv(
            temporary_paths[validation_csv],
            index=False,
            encoding="utf-8-sig",
        )

        with pd.ExcelWriter(
            temporary_paths[workbook_path], engine="openpyxl"
        ) as writer:
            summary_export.to_excel(
                writer,
                sheet_name="Validation_Summary",
                index=False,
            )
            detected_export.to_excel(
                writer,
                sheet_name="All_Exceptions",
                index=False,
            )
            comparison_export.to_excel(
                writer,
                sheet_name="Validation_Detail",
                index=False,
            )
            ground_truth_export.to_excel(
                writer,
                sheet_name="Ground_Truth",
                index=False,
            )
            row_counts_export.to_excel(
                writer,
                sheet_name="Data_Row_Counts",
                index=False,
            )

            for defect_type, sheet_name in SHEET_NAMES.items():
                category_df = detected_export[
                    detected_export["defect_type"] == defect_type
                ]
                category_df.to_excel(
                    writer,
                    sheet_name=sheet_name,
                    index=False,
                )

        style_workbook(temporary_paths[workbook_path])

        # Each destination is replaced atomically after all new files are prepared.
        # This avoids leaving truncated individual files if a write fails.
        for final_path, temporary_path in temporary_paths.items():
            os.replace(temporary_path, final_path)

    except PermissionError as exc:
        print(
            "Could not replace one of the report files. "
            "Close any open Excel workbook or CSV file in Excel, "
            "then run the script again. Existing completed reports were preserved "
            "where the operating system allowed it."
        )
        print(f"File detail: {exc}")
        return 1

    except OSError as exc:
        print(f"Could not save the reports: {exc}")
        return 1

    finally:
        for temporary_path in temporary_paths.values():
            try:
                temporary_path.unlink(missing_ok=True)
            except OSError:
                pass

    print("\nReconciliation report generated successfully.")
    print(f"Detected exceptions: {len(detected)}")
    print(f"Ground-truth exceptions: {len(ground_truth)}")
    print(f"Excel workbook: {workbook_path}")
    print(f"Exceptions CSV: {exception_csv}")
    print(f"Validation CSV: {validation_csv}")

    print("\nValidation summary:")
    print(summary.to_string(index=False))

    perfect_ids = bool(
        (summary["false_positives"] == 0).all()
        and (summary["false_negatives"] == 0).all()
    )
    perfect_amounts = bool((summary["amount_mismatches"] == 0).all())

    if perfect_ids and perfect_amounts:
        print(
            "\nPASS: Detected IDs and amounts match the ground truth "
            "for this synthetic dataset."
        )
        return 0

    print(
        "\nWARNING: Validation found identifier or amount mismatches. "
        "Check Validation_Detail in the workbook."
    )
    return 2


if __name__ == "__main__":
    raise SystemExit(main())
