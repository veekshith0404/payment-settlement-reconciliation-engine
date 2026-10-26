"""Smoke tests for generated-workbook formatting."""

from pathlib import Path
import tempfile
import unittest

from openpyxl import Workbook, load_workbook

from src.run_reconciliation import style_workbook


class ReportFormattingTests(unittest.TestCase):
    def test_currency_amounts_receive_two_decimal_number_format(self):
        with tempfile.TemporaryDirectory() as temporary_directory:
            path = Path(temporary_directory) / "sample.xlsx"
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.append(["defect_type", "amount_under_review"])
            worksheet.append(["MDR_OVERCHARGE", 5.5])
            workbook.save(path)

            style_workbook(path)

            styled = load_workbook(path, data_only=False)
            self.assertEqual(
                styled.active["B2"].number_format,
                "#,##0.00;[Red]-#,##0.00",
            )
            styled.close()


if __name__ == "__main__":
    unittest.main()
